#!/usr/bin/env python3
"""Build only synthetic fixtures. Requires python-docx, reportlab and Pillow."""
from pathlib import Path
import hashlib
import json
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image, ImageDraw, ImageFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer, PageBreak, Image as PDFImage
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parent
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "fixtures"
OUT.mkdir(parents=True, exist_ok=True)
(OUT / "assets").mkdir(exist_ok=True)
image = Image.new("RGB", (640, 220), "white")
d = ImageDraw.Draw(image)
font_path = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
bold_path = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
font = ImageFont.truetype(font_path, 18)
d.text((20, 5), "Bean height (cm)", fill="black", font=font)
d.rectangle((85, 122, 250, 180), fill="#337f91")
d.rectangle((360, 35, 525, 180), fill="#337f91")
d.text((130, 190), "Start: 4", fill="black", font=font)
d.text((405, 190), "End: 10", fill="black", font=font)
image.save(OUT / "assets/figure-001.png")

blocks = [
    {"kind":"heading","level":1,"text":"Garden measurement packet"},
    {"kind":"paragraph","text":"Student copy. Distribute only the student activity."},
    {"kind":"heading","level":2,"text":"Student activity"},
    {"kind":"paragraph","text":"Read Table 1, then use Figure 1 to compare the measurements."},
    {"kind":"paragraph","text":"Table 1. Garden observations"},
    {"kind":"table","rows":[["Plant","Start height (cm)","End height (cm)"],["Bean","4","10"],["Pea","6","9"]]},
    {"kind":"figure","asset":"assets/figure-001.png","alt":"Two bars labeled Start: 4 and End: 10, showing bean height in centimeters."},
    {"kind":"paragraph","text":"Figure 1. Bean height changes from 4 cm to 10 cm."},
    {"kind":"paragraph","text":"Q1. How many centimeters did the bean grow? Show the subtraction."},
    {"kind":"paragraph","text":"Q2. Which plant grew more? Use both differences."},
    {"kind":"paragraph","text":"Explain your reasoning."},
    {"kind":"heading","level":1,"text":"Teacher guide"},
    {"kind":"paragraph","text":"Teacher copy. Keep this complete answer key separate from student distribution."},
    {"kind":"heading","level":2,"text":"Answer key"},
    {"kind":"paragraph","text":"Table 2. Worked answers"},
    {"kind":"table","rows":[["Item","Worked answer","Common error"],["Q1","10 - 4 = 6 cm.","Subtracting in reverse order."],["Q2","Bean: 6 cm; pea: 3 cm.\nBean grew 3 cm more.","Comparing final heights instead of growth."]]},
    {"kind":"paragraph","text":"Rubric: 1 point for each correct difference; 1 point for comparison with units."},
    {"kind":"paragraph","text":"Explain your reasoning."},
]
doc = Document()
doc.sections[0].page_width = Inches(8.5)
doc.sections[0].page_height = Inches(11)
doc.styles["Normal"].font.name = "DejaVu Sans"
doc.styles["Normal"].font.size = Pt(11)
for style in ("Title", "Heading 1", "Heading 2"):
    doc.styles[style].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles[style].font.name = "DejaVu Sans"
    for border in list(doc.styles[style].element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
pdfmetrics.registerFont(TTFont("FixtureSans", font_path))
pdfmetrics.registerFont(TTFont("FixtureSans-Bold", bold_path))
styles = getSampleStyleSheet()
for style in styles.byName.values():
    style.fontName = "FixtureSans-Bold" if style.name.startswith("Heading") else "FixtureSans"
story = []
for i, b in enumerate(blocks, 1):
    b["id"] = "B" + str(i).zfill(4)
    page = 1 if i <= 11 else 2
    b["locator"] = "page/" + str(page) + "/block/" + str(i)
    if i == 12:
        story.append(PageBreak())
    if b["kind"] == "heading":
        para = doc.add_heading(b["text"], 0 if i == 1 else b["level"])
        if i == 12:
            para.paragraph_format.page_break_before = True
        story.append(Paragraph(escape(b["text"]), styles["Heading" + str(b["level"])]))
    elif b["kind"] == "paragraph":
        doc.add_paragraph(b["text"])
        story.append(Paragraph(escape(b["text"]), styles["BodyText"]))
    elif b["kind"] == "table":
        table = doc.add_table(rows=len(b["rows"]), cols=len(b["rows"][0]))
        table.style = "Table Grid"
        table.autofit = False
        widths = [1.1, 2.5, 2.4] if i == 16 else [2.0, 2.0, 2.0]
        for col, width in zip(table.columns, widths):
            col.width = Inches(width)
        for r, row in enumerate(b["rows"]):
            for c, text in enumerate(row):
                cell = table.cell(r, c)
                cell.text = text
                cell.width = Inches(widths[c])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr = cell._tc.get_or_add_tcPr()
                borders = OxmlElement("w:tcBorders")
                for side in ("top", "left", "bottom", "right"):
                    edge = OxmlElement("w:" + side)
                    edge.set(qn("w:val"), "single")
                    edge.set(qn("w:sz"), "4")
                    edge.set(qn("w:color"), "D9D9D9")
                    borders.append(edge)
                pr.append(borders)
                margins = OxmlElement("w:tcMar")
                for side in ("top", "left", "bottom", "right"):
                    edge = OxmlElement("w:" + side)
                    edge.set(qn("w:w"), "90")
                    edge.set(qn("w:type"), "dxa")
                    margins.append(edge)
                pr.append(margins)
                if r == 0:
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "E6EFF2")
                    pr.append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.space_before = Pt(3)
        pdf_rows = [[Paragraph(escape(c).replace("\n", "<br/>"), styles["BodyText"]) for c in r] for r in b["rows"]]
        t = Table(pdf_rows, colWidths=[95, 195, 190])
        t.setStyle(TableStyle([("GRID",(0,0),(-1,-1),0.5,colors.grey),("BACKGROUND",(0,0),(-1,0),colors.HexColor("#e6eff2")),("VALIGN",(0,0),(-1,-1),"TOP"),("BOTTOMPADDING",(0,0),(-1,-1),8)]))
        story.append(t)
    elif b["kind"] == "figure":
        p = doc.add_paragraph()
        inline = p.add_run().add_picture(str(OUT / b["asset"]), width=Inches(5))
        inline._inline.docPr.set("descr", b["alt"])
        b["sha256"] = hashlib.sha256((OUT / b["asset"]).read_bytes()).hexdigest()
        story.append(PDFImage(str(OUT / b["asset"]), width=384, height=132))
    story.append(Spacer(1, 8))
doc.save(OUT / "garden-packet.docx")
SimpleDocTemplate(str(OUT / "garden-packet.pdf"), pagesize=(612,792), topMargin=40, bottomMargin=40, leftMargin=60, rightMargin=60).build(story)
# The authoring inventory is NOT automatically a reviewed PDF inventory.
# A reviewer must inspect every rendered page before changing this status.
inv = {"schema_version":1,"source_name":"garden-packet.pdf","source_kind":"pdf",
       "coverage":"page-by-page visual inventory",
       "visual_inventory_review":{"status":"unreviewed","reviewer":"","page_count":2,"pages":[]},"blocks":blocks}
(OUT / "pdf-inventory.unreviewed.json").write_text(json.dumps(inv, indent=2) + "\n")
expected = {"text_blocks":15,"tables":2,"table_cells":18,"figures":1,"ordered_blocks":18,
            "questions":["Q1","Q2"],"answers":["10 - 4 = 6 cm.","Bean: 6 cm; pea: 3 cm.\nBean grew 3 cm more."],
            "repeated_paragraph":{"text":"Explain your reasoning.","count":2},
            "source":"newly authored synthetic fixture; not a real loss incident"}
(OUT / "expected.json").write_text(json.dumps(expected, indent=2) + "\n")
print(OUT)
